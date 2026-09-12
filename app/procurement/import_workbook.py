import sqlite3
import sys
from pathlib import Path

from openpyxl import load_workbook

SHEET = "TRADE LANES AND PORT COVERAGE"
DB_PATH = Path(__file__).resolve().parent / "procurement.db"

COLUMNS = {
    "TRADE LANE": "trade_lane",
    "Shipping Line Name": "shipping_line_name",
    "Type of Shipping Line": "shipping_line_type",
    "Major Countries": "major_countries",
    "Ports Covered": "ports_covered",
    "Rate getting system": "rate_getting_system",
    "Contact Person": "contact_person",
    "Contact Number": "contact_number",
    "Email Address": "email_address",
    "Remarks": "remarks",
}


def read_rows(path):
    workbook = load_workbook(path, data_only=False)

    try:
        sheet = workbook[SHEET]

        headers = [
            sheet.cell(2, col).value
            for col in range(1, 11)
        ]
        if headers != list(COLUMNS):
            raise ValueError(
                "Sheet headers differ from the expected format."
            )

        # Repeat the heading only within explicitly merged lane cells.
        merged_lanes = {}

        for area in sheet.merged_cells.ranges:
            if area.min_col == area.max_col == 1:
                for row in range(area.min_row, area.max_row + 1):
                    merged_lanes[row] = sheet.cell(
                        area.min_row, 1
                    ).value

        records = []

        for row in range(3, sheet.max_row + 1):
            values = [
                sheet.cell(row, col).value
                for col in range(1, 11)
            ]

            if not any(value is not None for value in values):
                continue

            if values[1] is None:
                if any(value is not None for value in values[2:]):
                    raise ValueError(
                        f"Row {row} has data but no shipping line."
                    )
                continue  # Trade-lane heading without a carrier.

            if any(
                sheet.cell(row, col).data_type == "f"
                for col in range(1, 11)
            ):
                raise ValueError(
                    f"Row {row} contains a formula; "
                    "review before importing."
                )

            values[0] = merged_lanes.get(row, values[0])

            # Preserve text, including phone numbers and long notes.
            values = [
                None if value is None else str(value)
                for value in values
            ]

            records.append(tuple(values) + (SHEET, row))

        if not records:
            raise ValueError("No procurement records found.")

        return records

    finally:
        workbook.close()


def main():
    if len(sys.argv) != 2:
        raise SystemExit(
            "Usage: python -m app.procurement.import_workbook "
            "WORKBOOK.xlsx"
        )

    records = read_rows(Path(sys.argv[1]))

    # Stop if the database already exists. Never overwrite it.
    with DB_PATH.open("xb"):
        pass

    connection = sqlite3.connect(DB_PATH)

    try:
        connection.execute("BEGIN")

        fields = ", ".join(
            f"{column} TEXT" for column in COLUMNS.values()
        )

        connection.execute(
            f"CREATE TABLE procurement_entries ("
            f"id INTEGER PRIMARY KEY, {fields}, "
            "source_sheet TEXT NOT NULL, "
            "source_row INTEGER NOT NULL, "
            "UNIQUE(source_sheet, source_row))"
        )

        names = ", ".join([
            *COLUMNS.values(),
            "source_sheet",
            "source_row",
        ])
        placeholders = ", ".join(["?"] * 12)

        connection.executemany(
            f"INSERT INTO procurement_entries ({names}) "
            f"VALUES ({placeholders})",
            records,
        )

        saved = connection.execute(
            f"SELECT {names} FROM procurement_entries "
            "ORDER BY source_row"
        ).fetchall()

        if saved != records:
            raise RuntimeError(
                "Imported values differ from the source."
            )

        connection.commit()

    except Exception:
        connection.rollback()
        raise

    finally:
        connection.close()

    print(f"Imported and checked {len(records)} procurement rows.")
    print(f"Database: {DB_PATH}")
    print("Contact Details sheet has not been imported.")


if __name__ == "__main__":
    main()