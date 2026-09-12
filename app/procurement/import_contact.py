import sqlite3
import sys
from pathlib import Path

from openpyxl import load_workbook

DB_PATH = Path(__file__).resolve().parent / "procurement.db"
SHEET = "Contact Details"

HEADERS = [
    "E13",
    "Booking contact / Agency",
    "Office Location",
    "Contact number",
    "Email ID",
    "Contact person name",
]


def text(value):
    return None if value is None else str(value)


def name_key(value):
    return " ".join(
        (value or "").replace("\xa0", " ").casefold().split()
    )


def read_contacts(path):
    workbook = load_workbook(path, data_only=False)

    try:
        sheet = workbook[SHEET]

        if [sheet.cell(1, c).value for c in range(1, 7)] != HEADERS:
            raise ValueError("Unexpected contact-sheet headers.")

        # Resolve explicitly merged cells.
        merged = {}
        for area in sheet.merged_cells.ranges:
            value = sheet.cell(area.min_row, area.min_col).value
            for row in range(area.min_row, area.max_row + 1):
                for col in range(area.min_col, area.max_col + 1):
                    merged[row, col] = value

        records = []
        current_shipping_line = None

        for row in range(2, sheet.max_row + 1):
            raw = [
                sheet.cell(row, col).value
                for col in range(1, 8)
            ]

            if not any(value is not None for value in raw):
                current_shipping_line = None
                continue

            if any(
                sheet.cell(row, col).value is not None
                for col in range(8, sheet.max_column + 1)
            ):
                raise ValueError(
                    f"Row {row} has extra columns; review before importing."
                )

            if any(
                sheet.cell(row, col).data_type == "f"
                for col in range(1, 8)
            ):
                raise ValueError(f"Formula found on contact row {row}.")

            values = [
                text(merged.get((row, col), raw[col - 1]))
                for col in range(1, 8)
            ]

            # In this workbook, blank carrier cells continue the block.
            if name_key(values[0]):
                current_shipping_line = values[0]

            values[0] = current_shipping_line

            # Do not infer an office location from an unmerged blank.
            records.append(tuple(values) + (SHEET, row))

        if not records:
            raise ValueError("No contact records found.")

        return records

    finally:
        workbook.close()


def main():
    if len(sys.argv) != 2:
        raise SystemExit(
            "Usage: python -m app.procurement.import_contacts WORKBOOK.xlsx"
        )

    records = read_contacts(Path(sys.argv[1]))

    # Open only an existing database.
    connection = sqlite3.connect(
        DB_PATH.as_uri() + "?mode=rw",
        uri=True,
    )

    try:
        connection.execute("BEGIN")

        existing = connection.execute("""
            SELECT name FROM sqlite_master
            WHERE type = 'table' AND name = 'shipping_line_contacts'
        """).fetchone()

        if existing:
            raise RuntimeError(
                "Contact table already exists. Stopped to avoid duplicates."
            )

        known_keys = {
            name_key(row[0])
            for row in connection.execute("""
                SELECT DISTINCT shipping_line_name
                FROM procurement_entries
            """)
            if name_key(row[0])
        }

        connection.execute("""
            CREATE TABLE shipping_line_contacts (
                id INTEGER PRIMARY KEY,
                shipping_line_name TEXT,
                booking_agency TEXT,
                office_location TEXT,
                contact_number TEXT,
                email_address TEXT,
                contact_person TEXT,
                additional_information TEXT,
                source_sheet TEXT NOT NULL,
                source_row INTEGER NOT NULL,
                procurement_match_key TEXT,
                UNIQUE(source_sheet, source_row)
            )
        """)

        imported = []
        unmatched = set()

        for record in records:
            key = name_key(record[0])
            match_key = key if key and key in known_keys else None

            if match_key is None:
                unmatched.add(record[0] or "[Missing shipping line]")

            imported.append(record + (match_key,))

        connection.executemany("""
            INSERT INTO shipping_line_contacts (
                shipping_line_name,
                booking_agency,
                office_location,
                contact_number,
                email_address,
                contact_person,
                additional_information,
                source_sheet,
                source_row,
                procurement_match_key
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, imported)

        saved = connection.execute("""
            SELECT shipping_line_name, booking_agency, office_location,
                   contact_number, email_address, contact_person,
                   additional_information, source_sheet, source_row,
                   procurement_match_key
            FROM shipping_line_contacts
            ORDER BY source_row
        """).fetchall()

        if saved != imported:
            raise RuntimeError("Saved contact values do not match the import.")

        connection.commit()

        matched = sum(record[-1] is not None for record in imported)
        print(f"Imported and checked {len(imported)} contact rows.")
        print(f"Rows matched by shipping-line name: {matched}")
        print(f"Rows needing name review: {len(imported) - matched}")

        if unmatched:
            print("\nUnmatched shipping-line labels:")
            for label in sorted(unmatched, key=str.casefold):
                print(f"- {label}")

    except Exception:
        connection.rollback()
        raise

    finally:
        connection.close()


if __name__ == "__main__":
    main()