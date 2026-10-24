"""Initial Shipment Sheet import. Default is a dry run; never reads Transport Sheet."""
import argparse
import hashlib
import json
from datetime import date,datetime
from pathlib import Path
from uuid import uuid4
from sqlalchemy import inspect, select, text, func
from app.shipments.fields import FIELDS,review_flags
from app.shipments.store import get_engine,metadata,shipments,imports
from app.shipments.service import record_history

def cell_text(value):
    if value is None:
        return None
    if isinstance(value,(date,datetime)):
        # Excel's stored calendar date, not an interpretation of nearby text.
        return value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
    if isinstance(value,float) and value.is_integer():
        return str(int(value))
    return str(value)

def read_workbook(path):
    import openpyxl
    workbook=openpyxl.load_workbook(path,read_only=True,data_only=False)
    cached=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        if 'SHIPMENT SHEET' not in workbook.sheetnames:
            raise ValueError('Expected SHIPMENT SHEET worksheet.')
        rows=workbook['SHIPMENT SHEET'].iter_rows(values_only=True)
        headers=next(rows)
        actual=[str(v or '').strip() for v in headers[:33]]
        if actual != list(FIELDS.values()):
            raise ValueError('Shipment headers changed. Review the mapping before importing.')
        result=[]
        cached_rows=cached['SHIPMENT SHEET'].iter_rows(min_row=2,values_only=True)
        for row_no,(values,cached_values) in enumerate(zip(rows,cached_rows),2):
            if not any(v is not None and str(v).strip() for v in values):
                continue
            if any(v is not None and str(v).strip() for v in values[33:]):
                raise ValueError(f'Row {row_no}: data outside the 33 mapped columns.')
            padded=list(values[:33])+[None]*max(0,33-len(values))
            notes=[]
            for index,value in enumerate(padded):
                if isinstance(value,str) and value.startswith('='):
                    notes.append({'field':list(FIELDS)[index], 'formula':value,
                                  'cached_value':cell_text(cached_values[index])})
                    padded[index]=cached_values[index]
            data=dict(zip(FIELDS,map(cell_text,padded)))
            if any(v is not None and len(v)>20000 for v in data.values()):
                raise ValueError(f'Row {row_no}: a cell exceeds the supported field length.')
            result.append({'source_row':row_no,'source_notes':notes,**data})
        if not result:
            raise ValueError('Shipment Sheet contains no data rows.')
        return result
    finally:
        workbook.close()
        cached.close()

def import_rows(rows,apply=False):
    digest=hashlib.sha256(json.dumps(rows,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
    engine=get_engine()
    with engine.begin() as conn:
        # Serializes initial imports; UI must remain stopped for initial setup.
        conn.execute(text("SELECT pg_advisory_xact_lock(86241021)"))
        exists=inspect(conn).has_table('shipments',schema='operations')
        if exists:
            if inspect(conn).has_table('shipment_imports',schema='operations'):
                if conn.scalar(select(imports.c.row_count).where(imports.c.checksum==digest)) is not None:
                    print('This exact Shipment Sheet was already imported; no changes made.')
                    return
            if conn.scalar(select(func.count()).select_from(shipments)):
                raise ValueError('Shipment table already contains records. Initial import refused to prevent duplicates or overwriting frontend edits.')
        if not apply:
            print(f'Dry run passed: {len(rows)} shipment rows. No database changes made.')
            return
        conn.execute(text('CREATE SCHEMA IF NOT EXISTS operations'))
        metadata.create_all(conn)
        for values in rows:
            row=conn.execute(shipments.insert().values(id=str(uuid4()),version=1,**values)
                .returning(shipments)).mappings().one()
            record_history(conn,row)
        if conn.scalar(select(func.count()).select_from(shipments)) != len(rows):
            raise ValueError('Row count verification failed; transaction will roll back.')
        saved=[{k:r[k] for k in rows[0]} for r in conn.execute(select(shipments).order_by(shipments.c.source_row)).mappings()]
        if saved != rows:
            raise ValueError('Imported values differ from source; transaction will roll back.')
        conn.execute(imports.insert().values(checksum=digest,row_count=len(rows)))
    print(f'Imported and verified {len(rows)} shipment rows. Procurement records unchanged.')

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook',type=Path)
    parser.add_argument('--apply',action='store_true')
    args=parser.parse_args()
    try:
        rows=read_workbook(args.workbook)
        flagged=sum(bool(review_flags(row)) for row in rows)
        print(f'Rows with date text needing review: {flagged}. Original values will be preserved.')
        import_rows(rows,args.apply)
    except ValueError as exc:
        raise SystemExit(str(exc)) from None
    except Exception as exc:
        raise SystemExit(f'Import failed: {type(exc).__name__}. Check the workbook path and PostgreSQL configuration. No connection details printed.') from None

if __name__=='__main__':
    main()
